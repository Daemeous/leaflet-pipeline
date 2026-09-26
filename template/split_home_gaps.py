"""
split_home_gaps.py  (TEMPLATE)
=================================
Cuts long homeless stretches out of roads that have them, so a leafletting
tracker's road list and map reflect where the houses actually are, not the
whole physical road.

Long rural/village roads often have their homes bunched together with a
long empty stretch between or around them (e.g. two ends of a village
strung along one road, or a road that runs past open fields before
reaching houses). Without this step the tracker treats the whole road as
one trackable unit: a leafletter has to walk/drive the empty stretch to
"finish" it, and marking one populated end Complete silently also claims
credit for an unvisited populated end further down the same road.

This mirrors route-planner's own trimToHomes() (leaflet-routes/js/graph.js)
-- same tuned constants (GAP_M / MARGIN_M / KEEP_IF_OVER_FRAC), so "empty"
means the same thing in the tracker as it does in route-planner -- but for
a different purpose. route-planner trims REMAINING geometry at routing
time to save driving/walking the empty bits and just drops them; this
script permanently splits the DATA sheet at build time so each home
cluster gets its own trackable Status, AND keeps the dropped stretches
(as Residences=0 rows) so a canvasser can see on the map that a stretch is
recorded-but-empty rather than not recorded at all.

Run this AFTER add_unnamed_roads.py (so its "Unknown Road" cluster rows --
never a named road split by a gap -- are left untouched) and BEFORE
build_tracker.py. No change was needed in build_tracker.py:
partition_route_only_rows() already pushes any Residences==0 row --
including the new gap rows this script adds -- below
ROUTE_PLANNER_MARKER, and every Dashboard/Checksum range is already
bounded to stop above that marker.

A road with 2+ home clusters far enough apart becomes "<Street> (Part 1)",
"<Street> (Part 2)", ... (residences split proportionally to each
cluster's matched UPRN count, rounded so the parts still sum to the
original total); a single-cluster road just keeps its name, trimmed to
the populated stretch. The empty stretches become new Residences=0 rows
named "<Street> [gap]" -- core.js renders these grey (vs. hot pink for a
genuinely all-empty road) under the existing route-only toggle.

A road that would split into more than MAX_PARTS clusters is left
unsplit instead. route-planner's own trimToHomes() has no reason to cap
this (a route just visits however many legs exist), but a tracker's road
list does -- tested against Stafford's real data, long radial roads with
miles of scattered ribbon development (Newport Road, Newcastle Road,
Eccleshall Road) split into 20+ tiny, barely-actionable parts each,
while the actually-intended case (a village-ish road with houses at two
or three points) landed at a median of 3.

Does its own UPRN buffer-match independently of estimate_residences_uprn.py
(same 40m buffer + commercial-cluster filter, but no far/orphan fallback --
this script only uses points to find where clusters are, not to compute
final counts, and any small imprecision from skipping the fallback washes
out because each split part's residence count is rescaled to the
*original* row's total, not recomputed from scratch).

CAVEAT: like the other steps that write <prefix>_Leafletting_residences_uprn.xlsx,
this changes row identity (row count, and split roads' Street names) for
every road it touches -- only safe on a fresh build, not a refresh with
live tracked progress, without the usual finalize_output.py merge care.

Requirements: geopandas, pandas, shapely, openpyxl, pyproj

Inputs:
    <prefix>_Leafletting_residences_uprn.xlsx  (from estimate_residences_uprn.py
                                                 + add_unnamed_roads.py)
    <prefix>_constituency.geojson
    ../data/osopenuprn_202605.csv

Output:
    <prefix>_Leafletting_residences_uprn.xlsx  (rewritten in place)
"""

import re
import time
import openpyxl
import pandas as pd
import geopandas as gpd
from shapely.geometry import LineString, MultiLineString
from shapely.ops import substring, linemerge

from constituency_config import get_config

_cfg = get_config()
_prefix = _cfg["output_prefix"]

TARGET_XLSX = f"{_prefix}_Leafletting_residences_uprn.xlsx"
CONSTITUENCY_GEOJSON = f"{_prefix}_constituency.geojson"
UPRN_CSV = "../data/osopenuprn_202605.csv"

ROAD_BUFFER_METRES = 40
COMMERCIAL_CLUSTER_THRESHOLD = 150

# Same tuned values as leaflet-routes/js/graph.js trimToHomes(), so "empty"
# agrees between the tracker and route-planner.
GAP_M = 250
MARGIN_M = 30
KEEP_IF_OVER_FRAC = 0.85
MIN_GAP_KEEP_M = 20   # don't split out a gap sliver shorter than this

# route-planner's trimToHomes() has no reason to cap cluster count -- a
# route just visits however many legs exist. A tracker's road list does:
# tested against Stafford's real data (2026-09-25), long radial roads with
# miles of scattered ribbon development (Newport Road, Newcastle Road,
# Eccleshall Road) split into 20+ tiny parts each, while the actually-
# intended case (a village-ish road with houses at two or three points)
# lands at a median of 3 parts. Beyond this cap, give up and leave the
# road as a single unsplit row rather than flooding the Data sheet with
# barely-actionable slivers.
MAX_PARTS = 6

GAP_SUFFIX = " [gap]"
ROUTE_PLANNER_MARKER = "###ROUTE_PLANNER_ONLY_BELOW###"


def parse_linestrings(geom_str):
    """Returns a list of shapely LineStrings, one per '|'-separated fragment."""
    if pd.isna(geom_str) or not str(geom_str).strip():
        return []
    out = []
    for part in str(geom_str).split("|"):
        m = re.match(r'LINESTRING\((.+)\)', part.strip())
        if not m:
            continue
        coords = []
        for pair in m.group(1).split(","):
            pts = pair.strip().split()
            if len(pts) == 2:
                try:
                    coords.append((float(pts[0]), float(pts[1])))
                except ValueError:
                    continue
        if len(coords) >= 2:
            out.append(LineString(coords))
    return out


def geoms_to_text(geoms):
    return "|".join(
        "LINESTRING(" + ", ".join(f"{x} {y}" for x, y in g.coords) + ")"
        for g in geoms if g is not None and len(g.coords) >= 2
    )


def cut_fragment(frag_4326, frag_27700, points_27700):
    """
    One fragment of one road. Returns (total_len, kept_stretches, gap_stretches)
    where kept_stretches/gap_stretches are lists of (geom_4326, homes_count)
    -- homes_count is 0 for gaps. Mirrors graph.js trimToHomes()'s per-
    fragment loop, but also keeps the complement (the gaps) instead of
    dropping it.
    """
    total = frag_27700.length
    if not points_27700 or total <= 0:
        return total, [], ([(frag_4326, 0)] if total >= MIN_GAP_KEEP_M else [])

    positions = sorted(frag_27700.project(p) for p in points_27700)

    clusters = []  # (start_pos, end_pos, count)
    i = 0
    while i < len(positions):
        j = i
        while j + 1 < len(positions) and positions[j + 1] - positions[j] <= GAP_M:
            j += 1
        clusters.append((positions[i], positions[j], j - i + 1))
        i = j + 1

    kept, gaps = [], []
    cursor = 0.0
    for start, end, count in clusters:
        a, b = max(0.0, start - MARGIN_M), min(total, end + MARGIN_M)
        if b - a <= 1:
            continue
        if a - cursor >= MIN_GAP_KEEP_M:
            gaps.append((substring(frag_4326, cursor / total, a / total, normalized=True), 0))
        kept.append((substring(frag_4326, a / total, b / total, normalized=True), count))
        cursor = b
    if total - cursor >= MIN_GAP_KEEP_M:
        gaps.append((substring(frag_4326, cursor / total, 1.0, normalized=True), 0))

    return total, kept, gaps


def merge_contiguous_fragments(frags_4326):
    """
    road_geometry stores one '|'-separated LINESTRING per OSM way segment --
    a single physical road is very often chopped into many of these at
    junctions (found via a real test run: a 2.1km road stored as ~30
    fragments). Stitch the ones that share an endpoint back into
    continuous lines first, so cluster/gap detection measures real
    distance along the road instead of spuriously treating every junction
    as a possible gap boundary. Segments that don't connect to anything
    (genuinely separate pieces) are left as their own entries.
    """
    if len(frags_4326) <= 1:
        return frags_4326
    merged = linemerge(MultiLineString(frags_4326))
    if merged.is_empty:
        return frags_4326
    return [merged] if merged.geom_type == "LineString" else list(merged.geoms)


def split_road(row, road_frags_4326, road_frags_27700, matched_points_27700):
    """
    Returns None if the road shouldn't be split (leave the row as-is), else
    a list of new row dicts (parts first, then gap rows) to replace it with.
    """
    # Each point belongs to whichever of the road's own fragments it's
    # actually nearest to, not every fragment it happens to be within the
    # buffer of (a fragment junction can put a point within range of two).
    pts_by_frag = [[] for _ in road_frags_27700]
    for p in matched_points_27700:
        nearest = min(range(len(road_frags_27700)), key=lambda k: road_frags_27700[k].distance(p))
        pts_by_frag[nearest].append(p)

    total_len, placed = 0.0, 0
    all_kept, all_gaps = [], []
    for f4326, f27700, pts_here in zip(road_frags_4326, road_frags_27700, pts_by_frag):
        t, kept, gaps = cut_fragment(f4326, f27700, pts_here)
        total_len += t
        all_kept.extend(kept)
        all_gaps.extend(gaps)
        placed += sum(c for _, c in kept)

    kept_len_m = sum(_geom_len_27700(g) for g, _ in all_kept)

    if not placed or not all_kept or kept_len_m >= KEEP_IF_OVER_FRAC * total_len:
        return None
    if len(all_kept) > MAX_PARTS:
        return None

    name = row["Street"]
    residences = int(row["Residences"] or 0)
    multi = len(all_kept) > 1
    new_rows = []
    running = 0
    for k, (geom, count) in enumerate(all_kept):
        share = residences if not multi else round(residences * count / placed)
        if multi and k == len(all_kept) - 1:
            share = residences - running
        running += share
        part_name = name if not multi else f"{name} (Part {k + 1})"
        new_rows.append(_row_dict(row, part_name, share, [geom]))

    for geom, _ in all_gaps:
        new_rows.append(_row_dict(row, f"{name}{GAP_SUFFIX}", 0, [geom]))

    return new_rows


def _geom_len_27700(geom_4326):
    return gpd.GeoSeries([geom_4326], crs="EPSG:4326").to_crs("EPSG:27700").iloc[0].length


def _row_dict(orig_row, street, residences, geoms_4326):
    merged = geoms_4326[0] if len(geoms_4326) == 1 else gpd.GeoSeries(geoms_4326, crs="EPSG:4326").unary_union
    centroid = merged.centroid
    return {
        "Street": street,
        "@lat": round(centroid.y, 6),
        "@lon": round(centroid.x, 6),
        "Ward": orig_row["Ward"],
        "Local Authority District": orig_row["Local Authority District"],
        "Status": "Not_Started",
        "Residences": residences,
        "road_geometry": geoms_to_text(geoms_4326),
        "partial_geometry": "-",
    }


def main():
    t0 = time.time()

    print("Step 1: Loading roads...")
    df = pd.read_excel(TARGET_XLSX, sheet_name="Data")
    df = df[df["Street"] != ROUTE_PLANNER_MARKER].reset_index(drop=True)

    print("Step 2: Loading constituency boundary & filtering UPRNs...")
    boundary = gpd.read_file(CONSTITUENCY_GEOJSON)
    constituency_boundary = boundary.geometry.iloc[0]
    constituency_padded = constituency_boundary.buffer(0.002)
    bbox = constituency_padded.bounds

    kept = []
    for chunk in pd.read_csv(
        UPRN_CSV, chunksize=500_000,
        usecols=["UPRN", "LATITUDE", "LONGITUDE"],
        dtype={"UPRN": "int64", "LATITUDE": "float64", "LONGITUDE": "float64"},
    ):
        pre = chunk[chunk["LONGITUDE"].between(bbox[0], bbox[2]) & chunk["LATITUDE"].between(bbox[1], bbox[3])]
        if len(pre):
            pts = gpd.GeoDataFrame(pre, geometry=gpd.points_from_xy(pre["LONGITUDE"], pre["LATITUDE"]), crs="EPSG:4326")
            kept.append(pts[pts.geometry.within(constituency_padded)][["UPRN", "geometry"]])
    uprn_gdf = pd.concat(kept, ignore_index=True)
    print(f"  {len(uprn_gdf):,} UPRNs in constituency")

    coord_key = uprn_gdf.geometry.y.round(5).astype(str) + ',' + uprn_gdf.geometry.x.round(5).astype(str)
    counts_at_coord = coord_key.map(coord_key.value_counts())
    uprn_gdf = uprn_gdf[counts_at_coord < COMMERCIAL_CLUSTER_THRESHOLD].copy()
    uprn_proj = uprn_gdf.to_crs("EPSG:27700")
    print(f"  {len(uprn_proj):,} after dropping commercial clusters")

    print("Step 3: Matching UPRNs to roads (40m buffer)...")
    frag_lists_4326 = [parse_linestrings(g) for g in df["road_geometry"]]
    road_idxs_with_geom = [i for i, frags in enumerate(frag_lists_4326) if frags]

    all_frags, frag_owner = [], []
    for i in road_idxs_with_geom:
        for f in frag_lists_4326[i]:
            all_frags.append(f)
            frag_owner.append(i)
    frags_gdf_27700 = gpd.GeoDataFrame({"road_idx": frag_owner}, geometry=all_frags, crs="EPSG:4326").to_crs("EPSG:27700")

    buffered = frags_gdf_27700.copy()
    buffered["geometry"] = frags_gdf_27700.buffer(ROAD_BUFFER_METRES)
    joined = gpd.sjoin(uprn_proj[["geometry"]], buffered[["road_idx", "geometry"]], how="inner", predicate="within")

    road_points = {}  # road_idx -> list of shapely Point (EPSG:27700)
    for uprn_i, group in joined.groupby(level=0):
        if len(group) == 1:
            ri = int(group["road_idx"].iloc[0])
        else:
            pt = uprn_proj.loc[uprn_i, "geometry"]
            ri = int(min(
                group["road_idx"],
                key=lambda ri: min(f.distance(pt) for f, o in zip(frags_gdf_27700.geometry, frag_owner) if o == ri)
            ))
        road_points.setdefault(ri, []).append(uprn_proj.loc[uprn_i, "geometry"])
    print(f"  {len(joined.index.unique()):,} UPRNs matched to {len(road_points)} roads")

    print("Step 4: Splitting roads with long homeless stretches...")
    n_split, n_parts, n_gaps = 0, 0, 0
    out_rows = []
    for i, row in df.iterrows():
        frags4326 = frag_lists_4326[i]
        residences = row.get("Residences")
        if not frags4326 or not isinstance(residences, (int, float)) or residences <= 0:
            out_rows.append(row.to_dict())
            continue
        pts = road_points.get(i, [])
        merged4326 = merge_contiguous_fragments(frags4326)
        merged27700 = gpd.GeoSeries(merged4326, crs="EPSG:4326").to_crs("EPSG:27700").tolist()
        new_rows = split_road(row, merged4326, merged27700, pts)
        if new_rows is None:
            out_rows.append(row.to_dict())
            continue
        n_split += 1
        for r in new_rows:
            out_rows.append(r)
            if r["Residences"] > 0:
                n_parts += 1
            else:
                n_gaps += 1

    print(f"  {n_split} roads split -> {n_parts} populated parts + {n_gaps} gap segments")

    print(f"\nStep 5: Writing '{TARGET_XLSX}'...")
    wb = openpyxl.load_workbook(TARGET_XLSX)
    ws = wb["Data"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

    ws.delete_rows(2, ws.max_row)
    for r_i, row_dict in enumerate(out_rows, start=2):
        for c_i, h in enumerate(headers, start=1):
            ws.cell(r_i, c_i).value = row_dict.get(h)
    wb.save(TARGET_XLSX)

    print(f"Complete in {time.time() - t0:.0f}s. Output: {TARGET_XLSX}")


if __name__ == "__main__":
    main()
